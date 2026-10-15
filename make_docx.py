# -*- coding: utf-8 -*-
"""生成《陇原非遗数字图谱》选题创意与方案策划书（DOCX）。"""
import json, os, sys, re
sys.stdout.reconfigure(encoding='utf-8')

ROOT = r'C:\Users\wqy\Documents\deepseek-harness\default-workspace\shumei-2026'
DATA = os.path.join(ROOT, 'prototype', 'js', 'data.js')

def load_data():
    txt = open(DATA, encoding='utf-8').read()
    m = re.search(r'window\.ICH_DATA\s*=\s*(\{.*\})\s*;', txt, re.S)
    return json.loads(m.group(1))

D = load_data()
items = D['items']; stats = D['stats']
from collections import Counter
by_city_c = Counter(it['city'] for it in items)
by_cat_c = Counter(it['category'] for it in items)
by_batch_c = Counter(it['batch'] for it in items)
by_city = stats['by_city']; by_cat = stats['by_category']
total = len(items); cities = len(by_city_c); cats = len(by_cat_c)
top_city = by_city_c.most_common(1)[0][0]; top_cat = by_cat_c.most_common(1)[0][0]
top_batch = by_batch_c.most_common(1)[0][0]
humanity = stats.get('humanity_list') or []
cutoff = stats.get('cutoff', '')

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = RGBColor(0x33, 0x24, 0x18)
PAPER = RGBColor(0xF9, 0xF4, 0xE8)
RED = RGBColor(0xB0, 0x39, 0x2C)
GREEN = RGBColor(0x4B, 0x8A, 0x63)
BLUE = RGBColor(0x2E, 0x5C, 0x8A)

doc = Document()
# 默认字体（含中文）
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(10.5)
style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

# A4 纵向，正文净宽 16.2cm
for sec in doc.sections:
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4); sec.right_margin = Cm(2.4)
TEXT_W_CM = 21.0 - 2.4 - 2.4   # 16.2cm

def set_cjk(run, name='Microsoft YaHei'):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn('w:eastAsia'), name)

def H(text, level=1, color=INK, size=None):
    p = doc.add_heading(level=level)
    r = p.add_run(text); set_cjk(r); r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    return p

def P(text='', align=None, bold=False, color=INK, size=10.5, italic=False):
    p = doc.add_paragraph()
    if align: p.alignment = align
    r = p.add_run(text); set_cjk(r); r.bold = bold; r.italic = italic
    r.font.color.rgb = color; r.font.size = Pt(size)
    return p

def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_lead:
        r = p.add_run(bold_lead); set_cjk(r); r.bold = True
        r.font.color.rgb = INK; r.font.size = Pt(10.5)
    r2 = p.add_run(text); set_cjk(r2); r2.font.size = Pt(10.5)
    return p

def numbered(text, bold_lead=None):
    p = doc.add_paragraph(style='List Number')
    if bold_lead:
        r = p.add_run(bold_lead); set_cjk(r); r.bold = True; r.font.color.rgb = INK; r.font.size = Pt(10.5)
    r2 = p.add_run(text); set_cjk(r2); r2.font.size = Pt(10.5)
    return p

def set_tcw(cell, w):
    tcPr = cell._tc.get_or_add_tcPr()
    el = tcPr.find(qn('w:tcW'))
    if el is None:
        el = OxmlElement('w:tcW'); tcPr.append(el)
    el.set(qn('w:w'), str(w)); el.set(qn('w:type'), 'dxa')

def table(headers, rows, widths=None, header_color=RED):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        widths = [TEXT_W_CM / len(headers)] * len(headers)
    dxa = [int(w / 2.54 * 1440) for w in widths]
    # 复用已有 tblW / tblLayout，避免重复元素导致布局回退
    tblPr = t._tbl.tblPr
    for tag, attrs in (('w:tblW', {'w:w': str(sum(dxa)), 'w:type': 'dxa'}),
                       ('w:tblLayout', {'w:type': 'fixed'})):
        el = tblPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); tblPr.append(el)
        for k, v in attrs.items():
            el.set(qn(k), v)
    # tblGrid 必须与列宽一致
    grid = t._tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for gc, w in zip(grid.findall(qn('w:gridCol')), dxa):
            gc.set(qn('w:w'), str(w))
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        rp = hdr[i].paragraphs[0].add_run(h); set_cjk(rp); rp.bold = True
        rp.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); rp.font.size = Pt(10)
        tcPr = hdr[i]._tc.get_or_add_tcPr()
        shd = tcPr.makeelement(qn('w:shd'), {qn('w:val'):'clear', qn('w:fill'):'B0392C'})
        tcPr.append(shd)
        set_tcw(hdr[i], dxa[i])
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            rr = cells[i].paragraphs[0].add_run(str(v)); set_cjk(rr); rr.font.size = Pt(9.5)
            set_tcw(cells[i], dxa[i])
    return t

# ============ 封面 ============
P('数字媒体科技作品及创意竞赛 · 数字艺术与可视化设计类', align=WD_ALIGN_PARAGRAPH.CENTER, color=RED, size=12, bold=True)
title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('《陇原非遗数字图谱》'); set_cjk(r); r.font.size = Pt(30); r.bold = True; r.font.color.rgb = INK
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('选题创意与方案策划书'); set_cjk(r); r.font.size = Pt(17); r.font.color.rgb = GREEN
meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run('交互式甘肃国家级非遗可视化系统'); set_cjk(r); r.font.size = Pt(12); r.italic = True; r.font.color.rgb = INK
doc.add_paragraph()
info = table(['项目', '内容'], [
    ['参赛类别', '数字艺术与可视化设计类（自主选题）'],
    ['参赛单位', '甘肃农业大学'],
    ['团队成员', '王乾勇、袁博宇'],
    ['数据基底', '甘肃省国家级非物质文化遗产代表性项目（五批名录，2006—2021）'],
    ['收录条目', f'{total} 条 · 覆盖 {cities}/14 市州 · {cats}/10 门类'],
    ['技术路线', '纯前端零依赖 HTML/CSS/JS；SVG + Canvas；离线可运行'],
    ['交付形态', '可运行原型 + 本策划书 + 答辩路演 PPT'],
], widths=[4, 12])
doc.add_page_break()

# ============ 一、项目背景 ============
H('一、项目背景与选题依据', 1)
P('非物质文化遗产是中华文明绵延传承的生动见证。甘肃省地处黄河上游、丝路咽喉，多民族共生、文化层积深厚，'
  '国家级非遗项目数量位居全国前列。然而，零散于各级政府名录中的项目信息多以表格、公文形式存在，'
  '缺乏一个能够"一览全省、纵览脉络、再生美学"的数字化入口。')
P('在"文化数字化"与"国家文化数字化战略"背景下，本作品以甘肃 14 市州、10 大门类的国家级非遗为数据基底，'
  '构建一套可交互、可检索、可再生的数字图谱，既回应非遗"活起来"的传播诉求，也契合数字媒体竞赛对'
  '"艺术性 × 技术性 × 文化价值"的综合考察。')
bullet('官方公布国家级非遗代表性项目 83 项（含扩展子项），省级 493 项，国家级代表性传承人 68 名，数据口径来自甘肃省文旅厅 2024 年公开表述。', '数据厚度：')
bullet('本图谱收录 76 条甘肃申报条目，覆盖全部 14 个市州与全部 10 个门类，五批名录齐全。', '收录完整：')
bullet('' + ('、'.join(humanity) if humanity else '花儿、环县道情皮影戏等') + ' 入选联合国教科文组织人类非遗名录，是甘肃非遗的"世界级名片"。', '国际影响：')

# ============ 二、创意来源 ============
H('二、创意来源与设计理念', 1)
P('作品以"图谱"为核心隐喻——既是地理之图（空间分布），也是关系之谱（门类—市州—批次—标签的关联网络），'
  '更是纹样之再生（从文化遗产中生长出新的视觉语言）。三条主线贯穿始终：')
numbered('以甘肃 14 市州行政边界为底图，按项目数量做色阶映射（示意分布，点位置于市州质心而非精确坐标），让观众"一眼看见"非遗的地理格局。', '空间图谱：')
numbered('用力导向图呈现"门类—项目—市州—标签"的关联谱系，点击任一项目即可展开简介、批次、传承人线索，把名录变成可漫游的知识网络。', '关系谱系：')
numbered('提取马家窑旋纹、敦煌藻井、香包、皮影、唐卡等纹样基因，以参数化方式"再生"出新的敦煌矿物色数字纹样，强调"致敬而非复制"的创作伦理。', '纹样再生：')
P('视觉系统取敦煌矿物颜料色谱（朱砂、石青、石绿、藤黄、赭石、紫檀）与宣纸／玄色双主题，'
  '整体气质"古朴而现代"，兼顾答辩现场的识别度与打印交付的清晰度。')

# ============ 三、需求分析 ============
H('三、需求与可行性分析', 1)
P('（一）用户与场景')
bullet('评审与答辩：需在投影/网页中快速呈现"数据体量 + 可视化美感 + 交互逻辑"。')
bullet('传播与科普：社会公众、文化研究者在浏览器中自助检索、了解甘肃非遗。')
bullet('教学与存档：作为课程设计/实践成果，附带可核查的数据来源说明。')
P('（二）可行性')
bullet('数据可得：国家级名录由政府网、中国非物质文化遗产网公开，已结构化整理为 76 条条目。', )
bullet('技术可行：纯前端方案无需服务器与第三方库，单文件即可离线运行，适配答辩现场网络不确定的情况。')
bullet('算力可行：Canvas 力导向与参数化纹样运算量极小，普通笔记本即可流畅运行。')

# ============ 四、总体方案 ============
H('四、总体方案与功能架构', 1)
P('作品为单页 Web 应用，包含五大视图，共用同一份数据与视觉系统：')
table(['视图', '功能', '技术实现'],
    [['总览', '数据速读、门类占比环图、市州排行、批次分布、关键事实', 'SVG + Canvas'],
     ['图谱地图', '甘肃区划 Choropleth、项目点位、点击下钻市州详情、按数量/门类丰富度切换', 'SVG 投影'],
     ['关系谱系', '门类—项目—市州—标签力导向网络，悬停高亮、拖拽、缩放、模式切换', 'Canvas 物理仿真'],
     ['纹样再生', '5 类纹样家族参数化生成，导出 2160×2160 PNG，宣纸/玄色主题', 'Canvas 程序化绘制'],
     ['关于', '选题缘起、数据来源、创作伦理与免责声明', 'HTML/CSS']],
    widths=[2.6, 8.4, 5])
P('（附）核心数据统计', bold=True, color=GREEN)
table(['指标', '数值'],
    [['收录条目', f'{total} 条'],
     ['覆盖市州', f'{cities} / 14（{top_city.replace("回族自治州","").replace("藏族自治州","")} 最多，{by_city_c[top_city]} 项）'],
     ['覆盖门类', f'{cats} / 10（{top_cat} 最多，{by_cat_c[top_cat]} 项）'],
     ['批次齐全', '第一批'+str(by_batch_c.get("第一批",0))+' · 第二批'+str(by_batch_c.get("第二批",0))+' · 第三批'+str(by_batch_c.get("第三批",0))+' · 第四批'+str(by_batch_c.get("第四批",0))+' · 第五批'+str(by_batch_c.get("第五批",0))+' 项'],
     ['官方口径', f'国家级 {stats.get("national_total","—")} 项 · 省级 {stats.get("provincial_total","—")} 项 · 传承人 {stats.get("inheritor_national","—")} 名'],
     ['数据截至', cutoff]],
    widths=[4, 12])

# ============ 五、技术路线 ============
H('五、技术路线与实现要点', 1)
bullet('零依赖离线运行：原生 HTML/CSS/JS，无框架、无 CDN、无构建步骤，双击 index.html 即可打开。', '工程：')
bullet('地理投影：自实现等距圆柱投影（按甘肃纬度余弦校正），将 DataV.GeoAtlas 边界数据绘制为 SVG。', '地图：')
bullet('力导向布局：自实现 Barnes-Hut 式松弛迭代（斥力 + 弹簧 + 向心力），预演算使首帧即稳定。', '网络：')
bullet('确定性随机：mulberry32 种子算法，保证点位散布与纹样演化可复现、可裁剪。', '算法：')
bullet('矿物色谱与双主题：敦煌色板 + 宣纸/玄色主题切换，全程支持打印。', '视觉：')

# ============ 六、数据与来源 ============
H('六、数据来源与真实性说明', 1)
P('项目数据来自公开权威渠道，并保留可核查来源：')
bullet('中国政府网（www.gov.cn）第一至第五批国家级非遗名录公布通知（含项目编号）。')
bullet('中国非物质文化遗产网（www.ihchina.cn）项目详情库。')
bullet('甘肃省文化和旅游厅公开名录与 2024 年非遗统计数据（经公开报道与转载交叉核对）。')
bullet('DataV.GeoAtlas 公开行政区划边界数据（仅几何，不含业务数据）。')
P('数据伦理与免责声明：', bold=True, color=RED)
bullet('地图点位为"示意分布"，统一置于市州质心并做确定性抖动，不代表项目精确坐标。')
bullet('纹样再生为参数化"致敬创作"，不复制任何文物的原貌与纹样，仅提取可公开的文化基因。')
bullet('个别条目编号依权威原文无法核实者留空，未做推测；统计口径差异已在 notes 中说明。')

# ============ 七、创新点 ============
H('七、创新点与特色', 1)
numbered('把"名录表格"转译为"可漫游的知识图谱"，空间、关系、美学三层叙事统一于一套视觉语言。')
numbered('纹样再生模块以算法"生长"出新的敦煌色数字纹样，提出非遗数字化的原创表达路径。')
numbered('纯前端零依赖、离线可运行，显著降低答辩与展演的工程风险，亦便于长年存档与传播。')
numbered('数据可核查、伦理有声明，体现数字媒体作品应有的学术规范与社会责任。')

# ============ 八、实施计划 ============
H('八、实施计划与进度', 1)
table(['阶段', '时间', '任务'],
    [['需求与数据', '已完成', '选题论证、名录结构化（76 条）、边界数据处理'],
     ['原型开发', '已完成', '五大视图、可视化引擎、双主题、导出'],
     ['打磨验证', '进行中', '多屏渲染校验、交互细节、性能优化'],
     ['文档与答辩', '进行中', '本策划书、路演 PPT、演示脚本'],
     ['提交参赛', '2026-10-23 前', '校赛报名与作品提交']],
    widths=[3.5, 3.5, 9])

# ============ 九、预期成效 ============
H('九、预期成效与价值', 1)
bullet('传播价值：为甘肃非遗提供兼具美感与信息密度的数字入口，利于科普与文旅传播。')
bullet('教学价值：作为数字媒体/可视化课程的完整实践案例，代码开源可读。')
bullet('竞赛价值：以"文化厚度 + 技术完成度 + 视觉表现"回应赛道评审要点。')

doc.add_paragraph()
foot = doc.add_paragraph(); foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = foot.add_run('甘肃农业大学 · 王乾勇、袁博宇 ｜ 2026 年第十四届全国大学生数字媒体科技作品及创意竞赛'); set_cjk(r)
r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x88, 0x80, 0x70)

out = os.path.join(ROOT, 'docs', '陇原非遗数字图谱_选题创意与方案策划书.docx')
doc.save(out)
print('saved', out, os.path.getsize(out), 'bytes')
