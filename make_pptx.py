# -*- coding: utf-8 -*-
"""生成《陇原非遗数字图谱》答辩/路演 PPTX。"""
import json, os, re, sys
from collections import Counter
sys.stdout.reconfigure(encoding='utf-8')

ROOT = r'C:\Users\wqy\Documents\deepseek-harness\default-workspace\shumei-2026'
DATA = os.path.join(ROOT, 'prototype', 'js', 'data.js')
IMG = os.path.join(ROOT, 'docs', 'images')

def load_data():
    txt = open(DATA, encoding='utf-8').read()
    return json.loads(re.search(r'window\.ICH_DATA\s*=\s*(\{.*\})\s*;', txt, re.S).group(1))

D = load_data(); items = D['items']; stats = D['stats']
by_city = Counter(it['city'] for it in items)
by_cat = Counter(it['category'] for it in items)
by_batch = Counter(it['batch'] for it in items)
total = len(items); cities = len(by_city); cats = len(by_cat)
top_city = by_city.most_common(1)[0]; top_cat = by_cat.most_common(1)[0]
humanity = stats.get('humanity_list') or []
cutoff = stats.get('cutoff', '')

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

INK   = RGBColor(0x33, 0x24, 0x18)
RED   = RGBColor(0xB0, 0x39, 0x2C)
GREEN = RGBColor(0x4B, 0x8A, 0x63)
BLUE  = RGBColor(0x2E, 0x5C, 0x8A)
GOLD  = RGBColor(0xC9, 0x9A, 0x3E)
PAPER = RGBColor(0xF9, 0xF4, 0xE8)
GREY  = RGBColor(0x8A, 0x80, 0x70)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height

def slide():
    return prs.slides.add_slide(BLANK)

def bg(s, color=PAPER):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color

def rect(s, x, y, w, h, color, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def txt(s, x, y, w, h, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, font='Microsoft YaHei', italic=False, line_spacing=1.06):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = Pt(2); tf.margin_right = Pt(2); tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    lines = text.split('\n')
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line_spacing
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color; r.font.name = font
        rPr = r._r.get_or_add_rPr(); rPr.set('{http://schemas.openxmlformats.org/drawingml/2006/main}lang', 'zh-CN')
    return tb

def bullets(s, x, y, w, h, items, size=16, color=INK, gap=6):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        lead, body = (it if isinstance(it, tuple) else (None, it))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.05
        if lead:
            r = p.add_run(); r.text = '◆ ' + lead + ('：' if not lead.endswith(('：', '，', '·')) else '')
            r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = RED; r.font.name = 'Microsoft YaHei'
        r2 = p.add_run(); r2.text = body
        r2.font.size = Pt(size); r2.font.color.rgb = color; r2.font.name = 'Microsoft YaHei'
    return tb

def header(s, kicker, title, num=None):
    rect(s, 0, 0, SW, Inches(1.18), INK)
    rect(s, 0, Inches(1.18), SW, Pt(3), RED)
    txt(s, Inches(0.6), Inches(0.16), Inches(11), Inches(0.34), kicker, size=12, color=GOLD, bold=True)
    txt(s, Inches(0.6), Inches(0.46), Inches(11.5), Inches(0.62), title, size=27, color=PAPER, bold=True)
    if num:
        txt(s, Inches(12.0), Inches(0.16), Inches(0.9), Inches(0.34), num, size=12, color=GREY, align=PP_ALIGN.RIGHT)

def pic_fit(s, path, x, y, maxw, maxh, align='center'):
    from PIL import Image
    iw, ih = Image.open(path).size
    ar = iw / ih
    w = maxw; h = int(w / ar)
    if h > maxh:
        h = maxh; w = int(h * ar)
    if align == 'center':
        px = x + (maxw - w) // 2
    elif align == 'left':
        px = x
    else:
        px = x + (maxw - w)
    py = y + (maxh - h) // 2
    s.shapes.add_picture(path, px, py, width=w, height=h)
    return px, py, w, h

# ---------- 1 封面 ----------
s = slide(); bg(s, INK)
rect(s, 0, Inches(2.5), SW, Inches(2.4), RGBColor(0x2A, 0x1E, 0x16))
rect(s, Inches(0.9), Inches(2.62), Inches(0.14), Inches(2.16), RED)
txt(s, Inches(1.2), Inches(2.55), Inches(11), Inches(1.1), '陇原非遗数字图谱', size=46, color=PAPER, bold=True)
txt(s, Inches(1.2), Inches(3.75), Inches(11), Inches(0.7), '交互式甘肃国家级非遗可视化系统', size=22, color=GOLD)
txt(s, Inches(1.2), Inches(4.55), Inches(11), Inches(0.5), '数字艺术与可视化设计类 · 自主选题', size=15, color=GREY)
txt(s, Inches(1.2), Inches(5.5), Inches(11), Inches(1.0),
    '甘肃农业大学\n王乾勇 · 袁博宇\n2026 第十四届全国大学生数字媒体科技作品及创意竞赛', size=15, color=PAPER, line_spacing=1.25)

# ---------- 2 选题背景 ----------
s = slide(); bg(s)
header(s, 'BACKGROUND', '为什么做：非遗的"数字入口"缺失', '02')
bullets(s, Inches(0.6), Inches(1.5), Inches(6.0), Inches(5.4), [
    ('背景：', '甘肃地处黄河上游、丝路咽喉，多民族共生，国家级非遗数量居全国前列。'),
    ('痛点：', '名录多以表格、公文形式存在，缺乏"一览全省、纵览脉络、再生美学"的数字化入口。'),
    ('机遇：', '"国家文化数字化战略"下，非遗"活起来"需要可交互、可检索、可再生的新型载体。'),
    ('定位：', '以 14 市州、10 大门类国家级非遗为基底，构建一套文化与技术兼具的数字图谱。'),
], size=15)
rect(s, Inches(6.9), Inches(1.6), Inches(5.9), Inches(5.0), RGBColor(0xEF, 0xE6, 0xD4))
bullets(s, Inches(7.2), Inches(1.85), Inches(5.3), Inches(4.6), [
    ('', f'官方公布国家级非遗 {stats.get("national_total","—")} 项（含扩展子项）'),
    ('', f'省级非遗 {stats.get("provincial_total","—")} 项'),
    ('', f'国家级代表性传承人 {stats.get("inheritor_national","—")} 名'),
    ('', f'本图谱收录 {total} 条甘肃申报条目'),
    ('', f'覆盖 {cities}/14 市州 · {cats}/10 门类'),
    ('', '五批名录齐全（2006—2021）'),
    ('', ('、'.join(humanity) if humanity else '花儿、环县道情皮影戏') + ' 入选人类非遗名录'),
], size=15, color=INK)

# ---------- 3 创意理念 ----------
s = slide(); bg(s)
header(s, 'CONCEPT', '做什么：图谱的三重叙事', '03')
cards = [
    ('空间图谱', '以甘肃区划为底图，按项目数做色阶映射，一眼看见非遗地理格局', RED),
    ('关系谱系', '力导向网络呈现"门类—项目—市州—标签"关联，点击即可漫游知识', GREEN),
    ('纹样再生', '参数化"生长"敦煌色数字纹样，致敬而非复制', BLUE),
]
cw = Inches(3.85); gap = Inches(0.36); x0 = Inches(0.75); y0 = Inches(1.7)
for i, (t, d, c) in enumerate(cards):
    x = x0 + i * (cw + gap)
    rect(s, x, y0, cw, Inches(4.4), PAPER, line=c)
    rect(s, x, y0, cw, Inches(0.12), c)
    txt(s, x, y0 + Inches(0.35), cw, Inches(0.7), t, size=24, color=c, bold=True, align=PP_ALIGN.CENTER)
    txt(s, x + Inches(0.25), y0 + Inches(1.3), cw - Inches(0.5), Inches(3.0), d, size=16, color=INK, align=PP_ALIGN.CENTER)
txt(s, Inches(0.75), Inches(6.35), Inches(11.8), Inches(0.7),
    '视觉系统：敦煌矿物色谱（朱砂/石青/石绿/藤黄/赭石/紫檀） + 宣纸·玄色双主题，古朴而现代。',
    size=14, color=GREY, align=PP_ALIGN.CENTER)

# ---------- 4 功能架构（加总览截图） ----------
s = slide(); bg(s)
header(s, 'PRODUCT', '产品总览：五大视图一体', '04')
pic_fit(s, os.path.join(IMG, 'shot-overview-crop.png'), Inches(0.45), Inches(1.42), Inches(8.3), Inches(5.5), 'center')
bullets(s, Inches(8.95), Inches(1.6), Inches(3.95), Inches(5.3), [
    ('总览', '数据速读 · 门类环图 · 市州排行 · 批次分布'),
    ('图谱地图', '甘肃区划 Choropleth + 项目点位下钻'),
    ('关系谱系', '力导向网络 · 悬停高亮 · 拖拽缩放'),
    ('纹样再生', '5 类纹样家族 · 导出 2160² PNG'),
    ('关于', '选题缘起 · 数据来源 · 创作伦理'),
], size=14)
txt(s, Inches(0.45), Inches(7.0), Inches(8.3), Inches(0.4), '▲ 总览视图：数据速读与四组统计图表', size=11, color=GREY, align=PP_ALIGN.CENTER)

# ---------- 5 功能细节：地图 + 谱系 ----------
s = slide(); bg(s)
header(s, 'FEATURES', '空间图谱 · 关系谱系', '05')
pic_fit(s, os.path.join(IMG, 'shot-map.png'), Inches(0.4), Inches(1.4), Inches(6.4), Inches(5.7), 'center')
pic_fit(s, os.path.join(IMG, 'shot-graph.png'), Inches(6.95), Inches(1.4), Inches(6.0), Inches(5.7), 'center')
txt(s, Inches(0.4), Inches(7.05), Inches(6.4), Inches(0.4), '▲ 图谱地图：色阶编码数量，与辖区面积无关', size=11, color=GREY, align=PP_ALIGN.CENTER)
txt(s, Inches(6.95), Inches(7.05), Inches(6.0), Inches(0.4), '▲ 关系谱系：124 节点 · 219 关联', size=11, color=GREY, align=PP_ALIGN.CENTER)

# ---------- 6 功能细节：纹样再生 ----------
s = slide(); bg(s, INK)
header(s, 'FEATURES', '纹样再生：从文化基因到数字纹样', '06')
pic_fit(s, os.path.join(IMG, 'shot-motif.png'), Inches(0.5), Inches(1.5), Inches(8.2), Inches(5.6), 'center')
bullets(s, Inches(9.0), Inches(1.7), Inches(3.9), Inches(5.2), [
    ('5 大家族', '马家窑旋纹 / 敦煌藻井 / 庆阳香包 / 皮影窗格 / 唐卡盘金'),
    ('参数驱动', '对称重数·层数·密度·配色实时可调，确定性随机可复现'),
    ('一键导出', '2160×2160 PNG，宣纸/玄色双主题，便于延展与印刷'),
    ('创作伦理', '提取文化基因做"致敬再生"，不复制文物原貌'),
], size=14, color=PAPER)
txt(s, Inches(0.5), Inches(7.1), Inches(8.2), Inches(0.4), '▲ 纹样再生视图（宣纸主题）', size=11, color=GREY, align=PP_ALIGN.CENTER)

# ---------- 7 数据来源 ----------
s = slide(); bg(s)
header(s, 'DATA', '数据来源与真实性', '07')
bullets(s, Inches(0.6), Inches(1.5), Inches(7.2), Inches(5.4), [
    ('权威渠道', '中国政府网五批国家级名录公布通知（含编号）；中国非物质文化遗产网项目库。'),
    ('交叉核对', '甘肃省文旅厅公开名录与 2024 年统计数据（经公开报道与转载交叉核对）。'),
    ('几何底图', 'DataV.GeoAtlas 公开行政区划边界（仅几何，不含业务数据）。'),
    ('伦理声明', '地图点位为示意分布（市州质心 + 确定性抖动），不代表精确坐标；纹样为参数化致敬创作。'),
    ('可核查', '编号无法核实者留空，未做推测；口径差异在文档 notes 中说明。'),
], size=15)
rect(s, Inches(8.1), Inches(1.6), Inches(4.7), Inches(5.0), RGBColor(0xEF, 0xE6, 0xD4))
bullets(s, Inches(8.35), Inches(1.85), Inches(4.2), Inches(4.6), [
    ('', f'收录条目：{total} 条'),
    ('', f'覆盖市州：{cities}/14（{top_city[0].replace("回族自治州","").replace("藏族自治州","")} {top_city[1]} 项）'),
    ('', f'覆盖门类：{cats}/10（{top_cat[0]} {top_cat[1]} 项）'),
    ('', '批次：一/二/三/四/五批齐全'),
    ('', f'截至：{cutoff}'),
], size=14, color=INK)

# ---------- 8 技术路线 ----------
s = slide(); bg(s)
header(s, 'TECH', '技术路线：纯前端 · 零依赖 · 离线', '08')
tech = [
    ('零依赖', '原生 HTML/CSS/JS，无框架、无 CDN、无构建，双击即开'),
    ('地理投影', '自实现等距圆柱投影（按纬度余弦校正），SVG 绘制区划'),
    ('力导向', '自实现斥力+弹簧+向心力松弛迭代，预演算使首帧即稳定'),
    ('确定性随机', 'mulberry32 种子算法，点位与纹样可复现、可裁剪'),
    ('敦煌色谱', '矿物色板 + 宣纸/玄色双主题，全程支持打印'),
]
bullets(s, Inches(0.6), Inches(1.5), Inches(12.0), Inches(5.4),
        [(t + '：', d) for t, d in tech], size=17, gap=12)

# ---------- 9 创新点 ----------
s = slide(); bg(s)
header(s, 'INNOVATION', '创新点与特色', '09')
bullets(s, Inches(0.6), Inches(1.5), Inches(12.0), Inches(5.4), [
    ('三层叙事', '把"名录表格"转译为可漫游的知识图谱，空间·关系·美学统一于一套视觉语言。'),
    ('原创表达', '纹样再生以算法"生长"新纹样，提出非遗数字化的原创路径。'),
    ('工程稳健', '纯前端零依赖、离线可运行，降低答辩与展演风险，便于长年存档传播。'),
    ('规范自觉', '数据可核查、伦理有声明，体现数字媒体作品的学术规范与社会责任。'),
], size=17, gap=12)

# ---------- 10 进度与成效 ----------
s = slide(); bg(s)
header(s, 'PLAN', '实施进度与预期价值', '10')
rect(s, Inches(0.6), Inches(1.5), Inches(6.0), Inches(5.2), PAPER, line=GREY)
txt(s, Inches(0.8), Inches(1.7), Inches(5.6), Inches(0.5), '实施进度', size=18, color=RED, bold=True)
table_data = [
    ('需求与数据', '已完成'), ('原型开发', '已完成'),
    ('打磨验证', '进行中'), ('文档与答辩', '进行中'),
    ('提交参赛', '2026-10-23 前'),
]
yy = Inches(2.4)
for t, v in table_data:
    txt(s, Inches(0.85), yy, Inches(3.2), Inches(0.5), t, size=15, color=INK)
    txt(s, Inches(4.0), yy, Inches(2.4), Inches(0.5), v, size=15, color=GREEN, bold=True)
    yy += Inches(0.78)
rect(s, Inches(6.9), Inches(1.5), Inches(5.9), Inches(5.2), PAPER, line=GREY)
txt(s, Inches(7.1), Inches(1.7), Inches(5.5), Inches(0.5), '预期价值', size=18, color=GREEN, bold=True)
bullets(s, Inches(7.15), Inches(2.35), Inches(5.4), Inches(4.2), [
    ('传播', '为甘肃非遗提供兼具美感与信息密度的数字入口'),
    ('教学', '数字媒体/可视化课程完整实践案例，代码可读'),
    ('竞赛', '以文化厚度+技术完成度+视觉表现回应评审要点'),
], size=15)

# ---------- 11 结尾 ----------
s = slide(); bg(s, INK)
rect(s, 0, Inches(3.05), SW, Pt(3), RED)
txt(s, Inches(0.8), Inches(2.15), Inches(11.7), Inches(0.9), '让陇原非遗，在数字图谱中“活”起来',
    size=32, color=PAPER, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(0.8), Inches(3.45), Inches(11.7), Inches(1.2),
    '甘肃农业大学 · 王乾勇 · 袁博宇\n2026 第十四届全国大学生数字媒体科技作品及创意竞赛',
    size=16, color=GOLD, align=PP_ALIGN.CENTER, line_spacing=1.4)

out = os.path.join(ROOT, 'docs', '陇原非遗数字图谱_答辩路演.pptx')
prs.save(out)
print('saved', out, os.path.getsize(out), 'bytes; slides', len(prs.slides._sldIdLst))
